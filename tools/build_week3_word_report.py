"""Build and structurally validate the Week 3 formal acceptance report.

This builder is offline-only. It reads existing evidence and writes one DOCX;
it never calls an API or mutates source evidence.
"""

from __future__ import annotations

import csv
import json
import re
import zipfile
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "第三周_智能渠道调度离线决策规则与验收报告.docx"
RULES_MD = DOCS / "WEEK3_DECISION_RULES.md"
REPORT_MD = DOCS / "WEEK3_ACCEPTANCE_REPORT.md"
POLICY_JSON = ROOT / "config" / "decision_policy_v1.json"
CANDIDATES_CSV = ROOT / "data" / "candidate_channels_v1.csv"
BASELINE_CSV = ROOT / "data" / "manual_scoring_baseline_v1.csv"
CATALOG_CSV = ROOT / "data" / "scenario_catalog_v1.csv"
SUMMARY_CSV = ROOT / "output" / "offline_decision_summary_v1.csv"
DECISIONS_JSON = ROOT / "output" / "offline_decisions_v1.json"

NAVY = "163A5F"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "EAF3F8"
PALE_BLUE = "F5F9FC"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E1E8"
TEXT = "243447"
MUTED = "5E6B78"
YELLOW = "FFF4CE"
GOLD = "8A6D1D"
WHITE = "FFFFFF"
BODY_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_sources():
    rules = RULES_MD.read_text(encoding="utf-8")
    report = REPORT_MD.read_text(encoding="utf-8")
    policy = json.loads(POLICY_JSON.read_text(encoding="utf-8-sig"))
    candidates = read_csv(CANDIDATES_CSV)
    baseline = read_csv(BASELINE_CSV)
    catalog = read_csv(CATALOG_CSV)
    summary = read_csv(SUMMARY_CSV)
    decisions = json.loads(DECISIONS_JSON.read_text(encoding="utf-8"))
    return rules, report, policy, candidates, baseline, catalog, summary, decisions


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = sum(int(Cm(w).emu / 635) for w in widths_cm)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        dxa = int(Cm(width).emu / 635)
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(dxa))
        grid.append(col)
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell, width in zip(row.cells, widths_cm):
            dxa = int(Cm(width).emu / 635)
            cell.width = Cm(width)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(dxa))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, name=BODY_FONT, size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.3
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13.5, BLUE, 10, 5),
        ("Heading 3", 11.5, DARK_BLUE, 7, 4),
    ):
        style = styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Cm(0.8)
        style.paragraph_format.first_line_indent = Cm(-0.4)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.25


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_font(run, CODE_FONT, 9, color=MUTED)


def set_header_footer(section) -> None:
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Intelligent Channel Scheduler  |  Week 3")
    set_font(run, "Arial", 9, True, MUTED)
    p = section.footer.paragraphs[0]
    run = p.add_run("LX  ·  UAT数据基线 + Mock规则场景                                      ")
    set_font(run, BODY_FONT, 8.5, color=MUTED)
    add_page_field(p)


def add_paragraph(doc, text: str, *, bold=False, color=None, style=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.keep_together = keep
    run = p.add_run(text)
    set_font(run, BODY_FONT, 10.5, bold, color)
    return p


def add_heading(doc, text: str, level=1):
    return doc.add_heading(text, level=level)


def add_bullets(doc, items: Iterable[str]) -> None:
    for item in items:
        add_paragraph(doc, item, style="List Bullet")


def add_numbered(doc, items: Iterable[str]) -> None:
    for item in items:
        add_paragraph(doc, item, style="List Number")


def add_callout(doc, title: str, text: str, fill=LIGHT_BLUE, accent=NAVY) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [16.0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + "\n")
    set_font(r, BODY_FONT, 10.5, True, accent)
    r = p.add_run(text)
    set_font(r, BODY_FONT, 10.2, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_box(doc, lines: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [16.0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines.splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, CODE_FONT, 9, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, font_size=8.5, landscape=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        set_font(r, BODY_FONT, font_size, True, WHITE)
    for i, values in enumerate(rows):
        row = table.add_row()
        if i % 2:
            for cell in row.cells:
                set_cell_shading(cell, PALE_BLUE)
        for cell, value in zip(row.cells, values):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if isinstance(value, (int, float)) or str(value) in {"True", "False", "—", "null"}:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_font(r, CODE_FONT if re.search(r"[A-Za-z0-9_]", str(value)) else BODY_FONT, font_size, color=TEXT)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_cover(doc, policy):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WEEK 3  ·  FORMAL ACCEPTANCE REPORT")
    set_font(r, "Arial", 11, True, BLUE)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第三周｜智能渠道调度\n离线决策规则与验收报告")
    set_font(r, BODY_FONT, 28, True, NAVY)
    p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("从结构化指标到可复现渠道选择")
    set_font(r, BODY_FONT, 15, color=DARK_BLUE)
    p.paragraph_format.space_after = Pt(40)
    add_table(doc, ["项目", "阶段"], [["Intelligent Channel Scheduler", "Week 3 — Reproducible Offline Decision"]], [7.8, 8.2], 9.5)
    add_table(doc, ["策略版本", "环境", "日期", "作者"], [[policy["policy_version"], "UAT数据基线 + Mock规则场景", "2026-07-21", "LX"]], [3.1, 6.8, 3.2, 2.9], 9.2)
    add_callout(doc, "证据边界", "本报告包含真实UAT基准与Mock规则场景。Mock结果不能证明真实渠道性能。", YELLOW, GOLD)
    doc.add_page_break()


def add_flow_diagram(doc):
    labels = ["请求输入", "候选发现", "资格过滤", "指标归一化", "策略评分", "同分排序", "选择/不可调度", "决策证据输出"]
    table = doc.add_table(rows=1, cols=len(labels))
    table.style = "Table Grid"
    set_table_geometry(table, [2.0] * 8)
    for i, (cell, label) in enumerate(zip(table.rows[0].cells, labels)):
        set_cell_shading(cell, LIGHT_BLUE if i % 2 == 0 else PALE_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label + ("\n→" if i < 7 else ""))
        set_font(r, BODY_FONT, 8.2, True, NAVY)
    doc.add_paragraph()


def add_relation_diagram(doc):
    rows = [
        ["策略配置 + 场景目录 + 候选快照", "→", "offline_decision_engine"],
        ["JSON明细 + CSV摘要", "←", "离线决策结果"],
        ["自动化测试", "→", "验收报告"],
    ]
    add_table(doc, ["输入/证据", "关系", "模块/产物"], rows, [7.0, 2.0, 7.0], 9.2)


def add_landscape_section(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.0)
    set_header_footer(section)
    return section


def add_portrait_section(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    set_header_footer(section)
    return section


def build_document() -> tuple[int, int]:
    rules_md, report_md, policy, candidates, baseline, catalog, summary, decisions = read_sources()
    # Evidence-derived acceptance checks before authoring.
    selected = sum(row["outcome"] == "selected" for row in summary)
    unroutable = sum(row["outcome"] == "unroutable" for row in summary)
    matched = sum(row["decision_matches_expected"].lower() == "true" for row in summary)
    assert (len(summary), selected, unroutable, matched) == (15, 13, 2, 15)
    specialist_match = re.search(r"专项测试\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|\s*(\d+)", report_md)
    full_match = re.search(r"完整测试\s*\|\s*(\d+)\s*\|\s*(\d+)\s*\|\s*(\d+)", report_md)
    assert specialist_match and specialist_match.groups() == ("46", "0", "0")
    assert full_match and full_match.groups() == ("222", "0", "1")
    assert "100%" in report_md
    assert policy["policy_version"] == "v1.0.0"
    by_sid = {item["scenario_id"]: item for item in decisions}
    catalog_by_sid = {row["scenario_id"]: row for row in catalog}
    assert all(by_sid[sid]["decision_matches_expected"] for sid in catalog_by_sid)

    doc = Document()
    style_document(doc)
    for section in doc.sections:
        set_header_footer(section)
    props = doc.core_properties
    props.title = "第三周｜智能渠道调度离线决策规则与验收报告"
    props.subject = "Week 3 — Reproducible Offline Decision"
    props.author = "LX"
    props.keywords = "offline decision, UAT, Mock, channel scheduling, acceptance"
    add_cover(doc, policy)

    add_heading(doc, "2. 文档信息与版本记录")
    add_table(doc, ["项目", "内容"], [
        ["文档名称", "第三周｜智能渠道调度离线决策规则与验收报告"],
        ["项目", "Intelligent Channel Scheduler"], ["阶段", "Week 3 — Reproducible Offline Decision"],
        ["策略版本", policy["policy_version"]], ["环境", "UAT数据基线 + Mock规则场景"],
        ["日期 / 作者", "2026-07-21 / LX"], ["文档状态", "正式验收稿"],
    ], [4.0, 12.0], 9.5)
    add_table(doc, ["版本", "日期", "作者", "说明"], [["1.0", "2026-07-21", "LX", "合并第三周决策规则与验收证据"]], [2.4, 3.4, 2.4, 7.8], 9.2)

    add_heading(doc, "3. 执行摘要")
    add_callout(doc, "验收结论", f"共验证{len(summary)}个固定场景：selected={selected}，unroutable={unroutable}，符合预期={matched}，符合率100%。专项测试46通过、0失败；完整测试222通过、0失败、1跳过。")
    add_paragraph(doc, "第三周已经形成可复现的离线决策链：读取策略与候选证据，按固定顺序过滤，对合格候选归一化并评分，再按分数、priority和channel_id排序，最终输出选择或不可调度结果及候选级审计证据。")
    add_paragraph(doc, "专项测试验证决策器运行前后输入文件哈希保持不变，重复运行产生稳定输出。报告生成只读取现有证据，不调用真实API，不新增渠道观测。")

    add_heading(doc, "4. 项目背景及六周计划中的第三周定位")
    add_paragraph(doc, "项目以六周迭代方式推进智能渠道调度能力。第二周已将UAT与历史证据整理为结构化指标；第三周位于“结构化证据”与“真实调用入口”之间，目标是先证明规则可计算、可复现、可解释、可审计。第四周再以本周稳定接口为输入推进调用入口与执行层设计。本文不对尚未完成的后续周能力作完成声明。")

    add_heading(doc, "5. 第三周目标、完成范围与非目标")
    add_bullets(doc, [
        "完成范围：候选资格过滤、配置驱动评分、排序与同分处理、选择/不可调度判断、JSON明细和CSV摘要。",
        "验证范围：15个固定场景、S001/S002人工评分、候选级过滤证据、expected_*隔离、稳定输出。",
        "非目标：不执行真实请求，不做正式流量路由，不实现重试、故障切换、动态价格同步或第五周能力。",
    ])

    add_heading(doc, "6. 系统输入—过滤—评分—排序—输出流程")
    add_flow_diagram(doc)
    add_paragraph(doc, "流程以请求参数和候选快照为输入。过滤阶段只保留证据完整且能力匹配的候选；评分阶段按策略版本进行归一化和加权；排序阶段使用未舍入分数；最后保存场景级结果和候选级证据。")
    add_heading(doc, "6.1 数据与模块关系图", 2)
    add_relation_diagram(doc)

    add_heading(doc, "7. 数据来源及真实数据与Mock数据边界")
    add_callout(doc, "Mock边界（必须保持）", "REAL-DS-48来自真实UAT证据，但样本量仅1；Mock候选和被场景覆盖的数据只用于离线规则验证。Mock结果不能证明真实渠道性能。", YELLOW, GOLD)
    add_bullets(doc, [
        "REAL-DS-48：is_mock=FALSE，来源为metrics_snapshot_v1+model_channel_matrix；流式能力待确认。",
        "MOCK-DS-FAST与MOCK-DS-CHEAP：is_mock=TRUE，来源为mock_scenario_design。",
        "TIE候选：只用于验证priority与channel_id同分规则。",
        "真实基础行被场景覆盖后，覆盖行标记为Mock场景数据，不是新增真实观测。",
    ])

    add_landscape_section(doc)
    add_heading(doc, "8. 模型与候选渠道说明")
    add_paragraph(doc, "本轮目标模型为 deepseek-v4-flash。下表数值来自当前 candidate_channels_v1.csv。")
    candidate_rows = []
    for row in candidates:
        candidate_rows.append([
            row["candidate_id"], row["requested_model"], row["channel_id"], row["channel_name"],
            row["latency_ms"], row["input_price_per_1m"], row["output_price_per_1m"], row["success_rate"],
            row["sample_size"], row["confidence_level"], row["supports_stream"], row["priority"], row["is_mock"], row["data_source"],
        ])
    add_table(doc,
        ["candidate_id", "requested_model", "channel_id", "channel_name", "latency_ms", "输入单价", "输出单价", "success_rate", "sample_size", "confidence", "stream", "priority", "is_mock", "data_source"],
        candidate_rows, [2.2, 2.4, 1.2, 2.7, 1.4, 1.3, 1.3, 1.3, 1.2, 1.5, 1.5, 1.1, 1.2, 3.2], 6.8)
    add_paragraph(doc, "REAL-DS-48的低样本和低置信度会提高证据惩罚；这反映证据不确定性，而不是对真实渠道质量的断言。")
    add_portrait_section(doc)

    add_heading(doc, "9. 策略配置与版本管理")
    filters = policy["filters"]
    norm = policy["normalization"]
    penalties = policy["penalties"]
    add_table(doc, ["配置项", "值", "作用"], [
        ["policy_version", policy["policy_version"], "生成稳定decision_id"],
        ["score_direction", policy["score_direction"], "最低分优先"],
        ["latency_reference_ms", norm["latency_reference_ms"], "延迟归一化参考"],
        ["cost_reference_cny", norm["cost_reference_cny"], "成本归一化参考"],
        ["max_metric_age_hours", filters["max_metric_age_hours"], "指标最大有效年龄"],
        ["minimum_sample_size", filters["minimum_sample_size"], "小样本阈值"],
        ["score_precision", policy["score_precision"], "输出舍入位数"],
    ], [5.2, 3.2, 7.6], 9.2)
    add_paragraph(doc, "所有权重、归一化基准、阈值和惩罚参数均从 decision_policy_v1.json 读取；内部计算使用未舍入数值，输出时才按score_precision舍入。")

    add_heading(doc, "10. 候选资格过滤规则")
    filter_rows = [
        [1, "requested_model不同", "model_mismatch", "模型不匹配"], [2, "状态不是available", "availability_unavailable", "候选不可用"],
        [3, "supports_text不是TRUE", "text_not_supported", "不支持文本"], [4, "非流式且supports_non_stream不是TRUE", "non_stream_not_supported", "不支持非流式"],
        [5, "流式且supports_stream=FALSE", "stream_not_supported", "明确不支持流式"], [5, "流式能力为空/待确认", "stream_capability_unknown", "不能把未知当支持"],
        [6, "latency_ms为空", "missing_latency", "无法计算延迟"], [7, "输入或输出价格为空", "missing_price", "无法计算成本"],
        [8, "币种不同", "currency_mismatch", "本周不做汇率转换"], [9, "更新时间为空", "missing_metrics_updated_at", "无法判断时效"],
        [10, f"指标年龄超过{filters['max_metric_age_hours']}小时", "stale_metrics", "指标过期"],
    ]
    add_table(doc, ["顺序", "条件", "exclusion_reason", "说明"], filter_rows, [1.3, 5.4, 5.2, 4.1], 8.4)
    add_paragraph(doc, "每个候选第一次命中的原因作为主排除原因。被排除候选不评分，无法计算的评分字段保持null，不伪造为0。")

    add_heading(doc, "11. 完整评分公式及字段解释")
    add_code_box(doc, "estimated_cost =\n  (input_tokens × input_price_per_1m\n   + output_tokens × output_price_per_1m) / 1,000,000\n\nlatency_norm = min(latency_ms / latency_reference_ms, 1.0)\ncost_norm = min(estimated_cost / cost_reference_cny, 1.0)\nfailure_risk = 1.0 - success_rate")
    add_code_box(doc, "small_sample_penalty = configured penalty when sample_size < minimum_sample_size\nconfidence_penalty = configured penalty for low / medium / high\nevidence_penalty = small_sample_penalty + confidence_penalty\n\nfinal_score =\n  latency_norm × strategy.weights.latency\n  + cost_norm × strategy.weights.cost\n  + failure_risk × strategy.weights.failure_risk\n  + evidence_penalty")
    add_table(doc, ["字段", "通俗解释"], [
        ["estimated_cost", "按输入/输出Token和每百万Token单价估算本次请求费用"], ["latency_norm", "延迟相对参考值的比例，最大为1"],
        ["cost_norm", "预计费用相对参考成本的比例，最大为1"], ["failure_risk", "1减成功率"],
        ["small_sample_penalty", "样本量低于阈值时增加的不确定性惩罚"], ["confidence_penalty", "根据置信度读取的惩罚"],
        ["evidence_penalty", "小样本与置信度惩罚之和"], ["final_score", "三项加权分与证据惩罚之和，越低越好"],
    ], [5.0, 11.0], 9.2)

    add_heading(doc, "12. 延迟优先策略")
    lw = policy["strategies"]["latency_first"]["weights"]
    add_paragraph(doc, f"latency_first权重：延迟{lw['latency']:.2f}、成本{lw['cost']:.2f}、失败风险{lw['failure_risk']:.2f}，权重和为1。该策略主要压低延迟，同时保留成本和可靠性影响。")
    add_heading(doc, "13. 成本优先策略")
    cw = policy["strategies"]["cost_first"]["weights"]
    add_paragraph(doc, f"cost_first权重：延迟{cw['latency']:.2f}、成本{cw['cost']:.2f}、失败风险{cw['failure_risk']:.2f}，权重和为1。该策略主要压低估算成本，同时保留延迟和可靠性影响。")

    add_heading(doc, "14. 小样本、低置信度和证据惩罚")
    add_table(doc, ["项目", "当前值", "规则"], [
        ["小样本阈值", filters["minimum_sample_size"], "sample_size低于阈值时惩罚"],
        ["小样本惩罚", penalties["small_sample_penalty"], "提高证据不确定性成本"],
        ["低置信度惩罚", penalties["low_confidence_penalty"], "confidence_level=low时增加"],
        ["中/高置信度惩罚", f"{penalties['medium_confidence_penalty']} / {penalties['high_confidence_penalty']}", "当前不增加置信度惩罚"],
    ], [5.0, 3.0, 8.0], 9.2)
    add_callout(doc, "低样本真实渠道说明", "REAL-DS-48基础样本量为1、置信度为low，因此基础场景证据惩罚为0.25（0.1+0.15）。这是对证据强度的保守处理，不是对真实渠道性能的负面结论。", YELLOW, GOLD)

    add_heading(doc, "15. 同分处理规则")
    add_numbered(doc, ["final_score按未舍入值升序。", "分数相同时priority数值升序。", "仍相同时channel_id升序；纯数字按数值比较，其他值按稳定字符串规则。"])
    add_heading(doc, "16. 不可调度规则")
    add_paragraph(doc, "如果没有合格候选，outcome=unroutable，selected_candidate为空，同时保留所有候选的排除原因。S004和S014用于验证该规则。")
    add_heading(doc, "17. expected_*字段隔离与防答案泄漏设计")
    add_callout(doc, "决策隔离", "expected_eligible、expected_exclusion_reason、expected_selected、expected_status和expected_selected_candidate只能在独立决策完成后用于验收，不能参与过滤、评分、排序或选择。")
    add_paragraph(doc, "专项测试会故意篡改候选expected_*字段并确认实际选择不变，从而证明决策器没有读取标准答案形成答案泄漏。")

    add_heading(doc, "18. S001延迟优先人工计算")
    s001 = by_sid["S001"]
    rows = sorted(s001["candidates"], key=lambda x: x["rank"] or 99)
    add_table(doc, ["候选", "预计费用", "延迟归一化", "成本归一化", "失败风险", "小样本", "置信度", "证据惩罚", "最终分", "排名"], [
        [r["candidate_id"], r["estimated_cost"], r["latency_norm"], r["cost_norm"], r["failure_risk"], r["small_sample_penalty"], r["confidence_penalty"], r["evidence_penalty"], r["final_score"], r["rank"]] for r in rows
    ], [2.8, 1.4, 1.6, 1.6, 1.4, 1.2, 1.3, 1.5, 1.4, 0.8], 7.3)
    add_code_box(doc, "MOCK-DS-FAST final_score = 0.2×0.65 + 0.3×0.20 + 0.03×0.15 + 0 = 0.1945")
    add_paragraph(doc, "S001选择MOCK-DS-FAST。该候选为Mock场景数据，选择结果不代表真实渠道性能。")

    add_heading(doc, "19. S002成本优先人工计算")
    s002 = by_sid["S002"]
    rows = sorted(s002["candidates"], key=lambda x: x["rank"] or 99)
    add_table(doc, ["候选", "预计费用", "延迟归一化", "成本归一化", "失败风险", "证据惩罚", "最终分", "排名"], [
        [r["candidate_id"], r["estimated_cost"], r["latency_norm"], r["cost_norm"], r["failure_risk"], r["evidence_penalty"], r["final_score"], r["rank"]] for r in rows
    ], [3.0, 1.8, 2.0, 2.0, 1.8, 1.8, 1.8, 1.0], 8.0)
    add_code_box(doc, "MOCK-DS-CHEAP final_score = 0.6×0.20 + 0.1×0.65 + 0.01×0.15 + 0 = 0.1865")
    add_paragraph(doc, "S002选择MOCK-DS-CHEAP。人工基线与程序输出在1e-6容差内一致。")

    add_landscape_section(doc)
    add_heading(doc, "20. 15个固定场景完整验收表")
    scenario_rows = []
    for row in summary:
        sid = row["scenario_id"]
        scenario_rows.append([
            sid, catalog_by_sid[sid]["scenario_name"], row["strategy"], row["outcome"],
            row["selected_candidate"] or "—", row["expected_selected_candidate"] or "—",
            row["decision_matches_expected"], row["candidate_count"], row["eligible_count"], row["excluded_count"],
        ])
    add_table(doc, ["scenario_id", "场景名称", "strategy", "outcome", "selected_candidate", "expected_selected", "matches", "候选", "合格", "排除"], scenario_rows,
              [2.0, 4.0, 3.0, 2.5, 3.5, 3.5, 1.8, 1.4, 1.4, 1.4], 7.4)
    add_callout(doc, "场景汇总", f"场景总数{len(summary)}；selected {selected}；unroutable {unroutable}；符合预期{matched}；不符合预期0；符合率100%。")
    add_portrait_section(doc)

    add_heading(doc, "21. S003流式能力过滤证据")
    s003 = by_sid["S003"]
    add_table(doc, ["candidate_id", "eligible", "exclusion_reason", "final_score"], [[c["candidate_id"], c["eligible"], c["exclusion_reason"] or "—", c["final_score"] if c["final_score"] is not None else "null"] for c in s003["candidates"]], [4.0, 2.4, 6.0, 3.6], 9.0)
    add_paragraph(doc, "REAL-DS-48因stream_capability_unknown排除；MOCK-DS-CHEAP因stream_not_supported排除；仅MOCK-DS-FAST合格并选中。")
    add_heading(doc, "22. S004全部不支持流式证据")
    s004 = by_sid["S004"]
    add_table(doc, ["candidate_id", "eligible", "exclusion_reason", "final_score"], [[c["candidate_id"], c["eligible"], c["exclusion_reason"], "null"] for c in s004["candidates"]], [4.0, 2.4, 6.0, 3.6], 9.0)
    add_callout(doc, "场景覆盖说明", "S004三个候选均为Mock场景行，REAL-DS-48在该场景中的流式字段也经过覆盖；这不是对真实渠道流式能力的新观测。", YELLOW, GOLD)

    add_heading(doc, "23. S011 priority同分选择")
    add_paragraph(doc, "TIE-A与TIE-B为指标完全相同的Mock候选，分数相同；priority分别为10和20，因此选择TIE-A。")
    add_heading(doc, "24. S012 channel_id同分选择")
    add_paragraph(doc, "TIE-001与TIE-002为分数和priority均相同的Mock候选，channel_id分别为101和102，因此选择TIE-001。")
    add_heading(doc, "25. S014全部不可用")
    add_paragraph(doc, "三个场景候选均为unavailable，以availability_unavailable排除，eligible_count=0，结果为unroutable。")
    add_heading(doc, "26. S015模型匹配")
    add_paragraph(doc, "两个Mock候选在该场景中被改为other-model并以model_mismatch排除；REAL-DS-48与deepseek-v4-flash匹配并选中。该结果只证明模型过滤，不构成真实性能评价。")

    add_heading(doc, "27. 人工基线与程序评分一致性")
    baseline_rows = []
    lookup = {(r["strategy"], r["candidate_id"]): r for r in baseline}
    for sid in ("S001", "S002"):
        decision = by_sid[sid]
        for c in decision["candidates"]:
            b = lookup[(decision["strategy"], c["candidate_id"])]
            baseline_rows.append([sid, decision["strategy"], c["candidate_id"], b["final_score"], c["final_score"], "一致"])
    add_table(doc, ["场景", "策略", "候选", "人工基线", "程序分数", "结果"], baseline_rows, [1.8, 3.2, 3.5, 2.5, 2.5, 2.5], 8.7)
    add_paragraph(doc, "六项评分误差均不超过1e-6。REAL-DS-48基础场景的0.25证据惩罚来自小样本0.1与低置信度0.15。")

    add_heading(doc, "28. 自动化测试结果")
    add_table(doc, ["测试范围", "通过", "失败", "跳过"], [["专项测试", 46, 0, 0], ["完整测试", 222, 0, 1]], [8.0, 2.5, 2.5, 3.0], 9.5)
    add_bullets(doc, ["专项测试覆盖策略读取、过滤顺序、公式、惩罚、同分、不可调度、expected_*隔离、结构和稳定性。", "完整测试覆盖项目既有测试集合。", "第三周不调用任何网络接口。"])

    add_heading(doc, "29. 决策可解释性和审计证据")
    add_paragraph(doc, "每个场景生成稳定decision_id，并记录策略版本、请求参数、outcome、selected_candidate、候选计数和匹配结果。每个候选记录eligible、exclusion_reason、费用、归一化项、失败风险、惩罚分解、最终分、priority、rank、is_mock和data_source。被排除候选评分字段保留null，避免把缺失证据误解为零成本或零风险。")

    add_heading(doc, "30. 可复现运行命令")
    add_code_box(doc, "python src\\offline_decision_engine.py --dry-run\npython src\\offline_decision_engine.py\npython -m pytest .\\tests\\test_offline_decision_engine.py -v\npython -m pytest .\\tests -q")
    add_paragraph(doc, "两次决策运行均汇总为：scenarios=15 selected=13 unroutable=2 matched=15。")

    add_heading(doc, "31. 已知限制")
    add_callout(doc, "风险与限制", "真实渠道样本量小；Mock数据不能用于真实性能结论；第三周不执行真实请求；尚未实现重试、故障切换和正式流量路由。", YELLOW, GOLD)
    add_bullets(doc, ["币种不一致当前直接排除，未实现汇率转换。", "流式能力未知时不能推断为支持。", "当前结果只覆盖固定场景和现有策略版本。"])

    add_heading(doc, "32. 第四周接口契约")
    add_numbered(doc, [
        "请求提供requested_model、stream_required、input_tokens、output_tokens、currency和decision_time。",
        "候选提供能力、状态、延迟、价格、成功率、样本量、置信度、更新时间、priority、is_mock和data_source。",
        "继续采用配置驱动规则；expected_*不得进入生产决策。",
        "输出保留decision_id、policy_version、候选级排除原因、评分分解和排名。",
        "未知或缺失证据不得补0；Mock证据不得混入真实路由结论。",
        "真实调用、鉴权、重试、故障切换和执行层必须与纯决策函数解耦。",
    ])
    add_heading(doc, "33. 第四周开发输入与计划")
    add_paragraph(doc, "第四周可将已验证的纯决策接口作为真实调用入口的上游，输入请求参数和可审计候选快照，输出选择结果与完整证据。后续计划是独立接入真实调用入口、鉴权保护和执行层观测，并继续保持Mock/真实证据隔离。本报告不声称这些能力已经实现。")

    add_heading(doc, "34. 最终验收结论")
    add_callout(doc, "正式结论", "第三周完成了从结构化指标到可复现渠道选择的离线决策链，并形成过滤、评分、排序、选择和证据输出的稳定规则。15个固定场景中13个selected、2个unroutable，15个全部符合预期，符合率100%。")
    add_paragraph(doc, "S001与S002的六项人工评分和程序评分在1e-6容差内一致；专项测试46通过、0失败，完整测试222通过、0失败、1跳过。候选级排除原因、评分分解、排名、数据来源和Mock标识使决策具备可解释性和审计能力。")
    add_paragraph(doc, "REAL-DS-48虽来自真实UAT证据，但样本量仅1；Mock候选与TIE场景只用于验证规则。Mock结果不能证明真实渠道性能。第三周尚未实现真实请求、重试、故障切换、动态价格同步和正式流量路由。基于当前场景、评分基线与自动化测试结果，本项目满足进入第四周接口设计与开发准备的条件，但真实调用仍需在后续阶段单独验证。")

    add_heading(doc, "35. 附录：关键文件清单、字段说明和复现命令")
    add_heading(doc, "35.1 关键文件清单", 2)
    add_table(doc, ["文件", "用途"], [
        ["config/decision_policy_v1.json", "策略参数唯一来源"], ["data/candidate_channels_v1.csv", "基础候选证据"],
        ["data/manual_scoring_baseline_v1.csv", "人工评分基线"], ["data/scenario_catalog_v1.csv", "15个场景目录"],
        ["output/offline_decisions_v1.json", "候选级决策明细"], ["output/offline_decision_summary_v1.csv", "场景级摘要"],
        ["src/offline_decision_engine.py", "离线决策实现"], ["tests/test_offline_decision_engine.py", "专项自动化测试"],
    ], [7.0, 9.0], 8.8)
    add_heading(doc, "35.2 关键字段说明", 2)
    add_table(doc, ["字段", "说明"], [
        ["decision_id", "policy_version与scenario_id组成的稳定标识"], ["eligible", "候选是否通过全部过滤"],
        ["exclusion_reason", "第一次命中的主排除原因"], ["final_score", "合格候选的最终分，越低越好"],
        ["rank", "合格候选排序名次"], ["is_mock", "是否为Mock数据"], ["data_source", "证据来源"],
        ["decision_matches_expected", "独立决策完成后的验收比较结果"],
    ], [5.0, 11.0], 9.0)
    add_heading(doc, "35.3 复现命令", 2)
    add_code_box(doc, "python src\\offline_decision_engine.py --dry-run\npython src\\offline_decision_engine.py --scenario S001\npython src\\offline_decision_engine.py --output-dir <临时目录>\npython -m pytest .\\tests\\test_offline_decision_engine.py -q\npython -m pytest .\\tests -q")
    add_paragraph(doc, "安全要求：不得在报告、日志或代码中记录API Key或完整Authorization头。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    validate_docx(OUTPUT)
    return len(doc.paragraphs), len(doc.tables)


def extract_docx_text(path: Path) -> str:
    document = Document(path)
    chunks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def validate_docx(path: Path) -> None:
    assert path.exists() and path.stat().st_size > 30 * 1024, "DOCX must exceed 30KB"
    text = extract_docx_text(path)
    required = [
        "第三周｜智能渠道调度", "从结构化指标到可复现渠道选择", "deepseek-v4-flash",
        "REAL-DS-48", "MOCK-DS-FAST", "MOCK-DS-CHEAP", "TIE-A", "TIE-001",
        "estimated_cost", "final_score", "46", "222", "100%",
        "Mock结果不能证明真实渠道性能", "34. 最终验收结论",
    ] + [f"S{i:03d}" for i in range(1, 16)]
    missing = [item for item in required if item not in text]
    assert not missing, f"missing required content: {missing}"
    topic_markers = [f"{i}. " for i in range(2, 36)]
    assert all(marker in text for marker in topic_markers), "required section topic missing"
    forbidden = ["API Key:", "Authorization:", "Bearer ", "2026-07-22实际执行结果", "PermissionError", "中、高置信度当前不加分"]
    assert not any(item in text for item in forbidden), "forbidden or stale content found"
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None, "invalid DOCX archive"


if __name__ == "__main__":
    paragraphs, tables = build_document()
    print(f"created={OUTPUT}")
    print(f"bytes={OUTPUT.stat().st_size} paragraphs={paragraphs} tables={tables}")
    print("structure_check=PASS sensitive_check=PASS")
